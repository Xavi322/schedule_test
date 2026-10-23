import os
import re
import json
import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

def ensure_template(template_path):
    """Создает template.html с выбором ПОТОКА (Очное/ИУП), НЕДЕЛИ и ГРУППЫ."""
    lines = [
        '<!DOCTYPE html><html lang="ru"><head><meta charset="UTF-8">',
        '<meta name="viewport" content="width=device-width, initial-scale=1.0">',
        '<meta name="theme-color" content="#1e293b">',
        '<link rel="manifest" href="manifest.json">',
        '<link rel="icon" type="image/png" sizes="192x192" href="app-icon-192.png">',
        '<link rel="icon" type="image/png" sizes="512x512" href="app-icon-512.png">',
        '<link rel="apple-touch-icon" sizes="192x192" href="app-icon-192.png">',
        '<title>Расписание занятий</title><style>',
        ':root{--bg:#f4f6f9;--card:#fff;--pri:#3b82f6;--pri-l:#eff6ff;--pur:#7c3aed;--pur-l:#f5f3ff;--pur-b:#ddd6fe;--warn:#ea580c;--warn-l:#fff7ed;--warn-b:#fed7aa;--dng:#dc2626;--dng-l:#fef2f2;--dng-b:#fecaca;--txt:#1e293b;--mut:#64748b;--brd:#e2e8f0;--grn:#10b981}',
        '*{box-sizing:border-box;margin:0;padding:0;font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,sans-serif}',
        'body{background:var(--bg);color:var(--txt);padding:16px}',
        'header{display:flex;justify-content:space-between;align-items:center;flex-wrap:wrap;gap:12px;margin-bottom:16px;padding:12px 16px;background:var(--card);border-radius:12px;box-shadow:0 1px 3px rgba(0,0,0,.05)}',
        'header h1{font-size:1.2rem;font-weight:700}',
        '.controls{display:flex;align-items:center;gap:8px;flex-wrap:wrap}',
        '.mode-select{background:#475569;color:#fff;padding:8px 12px;border-radius:8px;font-weight:600;font-size:.9rem;border:none;cursor:pointer;outline:none}',
        '.mode-select option{background:#fff;color:var(--txt)}',
        '.week-select{background:#1e293b;color:#fff;padding:8px 12px;border-radius:8px;font-weight:600;font-size:.9rem;border:none;cursor:pointer;outline:none}',
        '.week-select option{background:#fff;color:var(--txt)}',
        '.group-select{background:var(--pri);color:#fff;padding:8px 14px;border-radius:8px;font-weight:600;font-size:.95rem;border:none;cursor:pointer;outline:none}',
        '.group-select option{background:#fff;color:var(--txt)}',
        '.save-btn{background:#e2e8f0;color:var(--txt);padding:8px 12px;border-radius:8px;font-weight:600;font-size:.85rem;border:none;cursor:pointer;transition:.2s}',
        '.save-btn.is-saved{background:var(--grn);color:#fff}',
        '.week-grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(220px,1fr));gap:12px;align-items:start}',
        '.day-column{background:var(--card);border-radius:12px;border:1px solid var(--brd);overflow:hidden;box-shadow:0 1px 2px rgba(0,0,0,.04)}',
        '.day-header{background:#1e293b;color:#fff;padding:10px 12px;font-weight:600;font-size:.95rem;text-align:center}',
        '.day-slots{padding:8px;display:flex;flex-direction:column;gap:8px}',
        '.slot-card{border-radius:8px;padding:10px;min-height:92px;display:flex;flex-direction:column;justify-content:space-between}',
        '.active-slot{background:var(--pri-l);border:1px solid #bfdbfe;border-left:4px solid var(--pri)}',
        '.moved-slot{background:var(--pur-l);border:1px solid var(--pur-b);border-left:4px solid var(--pur)}',
        '.shifted-slot{background:var(--warn-l);border:1px solid var(--warn-b);border-left:4px solid var(--warn)}',
        '.empty-slot{background:#f8fafc;border:1px dashed #cbd5e1;color:#94a3b8;min-height:60px}',
        '.slot-header{display:flex;justify-content:space-between;align-items:center;font-size:.75rem;font-weight:600;color:var(--mut);margin-bottom:6px}',
        '.old-time{text-decoration:line-through;color:#94a3b8;font-size:.7rem;margin-right:4px}',
        '.custom-time{background:var(--warn);color:#fff;padding:2px 6px;border-radius:4px;font-weight:700;font-size:.75rem}',
        '.subject-title{font-size:.85rem;font-weight:600;line-height:1.35;color:var(--txt);margin-bottom:8px;white-space:pre-line}',
        '.empty-label{font-size:.8rem;text-align:center;margin:auto 0}',
        '.slot-footer{display:flex;justify-content:space-between;align-items:center;gap:4px;flex-wrap:wrap}',
        '.room-badge{background:#fff;border:1px solid #cbd5e1;color:#0f172a;font-size:.75rem;font-weight:600;padding:2px 8px;border-radius:6px;margin-left:auto}',
        '.moved-badge{background:var(--pur);color:#fff;font-size:.7rem;font-weight:600;padding:2px 6px;border-radius:6px}',
        '.online-btn{background:var(--grn);color:#fff;border:none;font-size:.75rem;font-weight:600;padding:4px 8px;border-radius:6px;cursor:pointer}',
        '.day-off{padding:32px 12px;text-align:center;color:var(--mut);font-weight:500}',
        '.cancelled-section{margin:0 8px 8px 8px;padding:8px 10px;background:var(--dng-l);border:1px solid var(--dng-b);border-radius:8px;font-size:.75rem;color:#7f1d1d}',
        '.cancelled-title{font-weight:700;color:var(--dng);margin-bottom:6px}',
        '.cancelled-item{margin-top:5px;line-height:1.35}',
        '.reason-tag{display:inline-block;background:#fff;border:1px solid var(--dng-b);color:var(--dng);font-weight:700;font-size:.68rem;padding:1px 5px;border-radius:4px;margin-right:3px}',
        '.modal-overlay{display:none;position:fixed;inset:0;background:rgba(15,23,42,.6);align-items:center;justify-content:center;z-index:1000;padding:16px}',
        '.modal-box{background:#fff;padding:20px;border-radius:12px;max-width:420px;width:100%;box-shadow:0 10px 25px rgba(0,0,0,.15)}',
        '.modal-title{font-weight:700;font-size:1rem;margin-bottom:10px}',
        '.modal-link{word-break:break-all;background:#f1f5f9;padding:10px;border-radius:8px;font-size:.85rem;margin-bottom:14px;border:1px solid var(--brd)}',
        '.modal-link a{color:var(--pri);font-weight:600;text-decoration:none}',
        '.modal-actions{display:flex;gap:8px;justify-content:flex-end}',
        '.btn{padding:8px 14px;border-radius:8px;font-weight:600;font-size:.85rem;cursor:pointer;border:none}',
        '.btn-copy{background:var(--pri);color:#fff}',
        '.btn-close{background:#e2e8f0;color:var(--txt)}',
        '</style></head><body>',
        '<header><h1>📅 Расписание</h1>',
        '<div class="controls">',
        '<select id="modeSelect" class="mode-select" autocomplete="off"></select>',
        '<select id="weekSelect" class="week-select" autocomplete="off"></select>',
        '<select id="groupSelect" class="group-select" autocomplete="off"></select>',
        '<button id="saveGroupBtn" class="save-btn" onclick="saveMyGroup()">📌 Запомнить</button>',
        '</div></header>',
        '<div id="weekGrid" class="week-grid"></div>',
        '<div id="linkModal" class="modal-overlay" onclick="if(event.target===this)closeModal()">',
        '<div class="modal-box"><div class="modal-title">💻 Ссылка на подключение</div>',
        '<div class="modal-link"><a id="modalLinkAnchor" href="#" target="_blank"></a></div>',
        '<div class="modal-actions">',
        '<button class="btn btn-copy" onclick="copyModalLink()">Скопировать</button>',
        '<button class="btn btn-close" onclick="closeModal()">Закрыть</button>',
        '</div></div></div><script>',
        'var appData = __APP_DATA_JSON__;',
        'function showLinkModal(u){var a=document.getElementById("modalLinkAnchor");a.href=u;a.textContent=u;document.getElementById("linkModal").style.display="flex";}',
        'function closeModal(){document.getElementById("linkModal").style.display="none";}',
        'function copyModalLink(){navigator.clipboard.writeText(document.getElementById("modalLinkAnchor").textContent);alert("Ссылка скопирована!");}',
        'var modeEl=document.getElementById("modeSelect");',
        'var weekEl=document.getElementById("weekSelect");',
        'var selectEl=document.getElementById("groupSelect");',
        'var saveBtn=document.getElementById("saveGroupBtn");',
        'var grid=document.getElementById("weekGrid");',
        'var STORAGE_KEY="pinned_group_v4";',
        'var modeKeys=Object.keys(appData);',
        'modeKeys.forEach(function(k){var o=document.createElement("option");o.value=k;o.textContent=appData[k].title;modeEl.appendChild(o);});',
        'if(modeKeys.length<=1){modeEl.style.display="none";}',
        'function getSavedState(){try{return localStorage.getItem(STORAGE_KEY)||window.name||"";}catch(e){return "";}}',
        'function populateMode(modeKey,prefGrp){',
        '  var mObj=appData[modeKey]||appData[modeKeys.slice().shift()];',
        '  if(!mObj)return;',
        '  var weeksPayload=mObj.weeks||[];',
        '  weekEl.innerHTML="";var currIdx=0;',
        '  weeksPayload.forEach(function(w,idx){var o=document.createElement("option");o.value=idx;o.textContent=w.label;if(w.is_current)currIdx=idx;weekEl.appendChild(o);});',
        '  weekEl.value=currIdx;weekEl.style.display=(weeksPayload.length<=1)?"none":"";',
        '  var firstW=weeksPayload.slice().shift()||{groups:{}};',
        '  var grps=Object.keys(firstW.groups||{});selectEl.innerHTML="";',
        '  grps.forEach(function(g){var o=document.createElement("option");o.value=g;o.textContent=g;selectEl.appendChild(o);});',
        '  if(prefGrp&&firstW.groups&&firstW.groups[prefGrp]){selectEl.value=prefGrp;}',
        '}',
        'function updateSaveBtn(){',
        '  var cur=modeEl.value+"|"+selectEl.value;',
        '  if(getSavedState()===cur){saveBtn.textContent="✅ Моя группа";saveBtn.className="save-btn is-saved";}',
        '  else{saveBtn.textContent="📌 Запомнить";saveBtn.className="save-btn";}',
        '}',
        'function saveMyGroup(){',
        '  var val=modeEl.value+"|"+selectEl.value;',
        '  try{localStorage.setItem(STORAGE_KEY,val);window.name=val;}catch(e){}',
        '  updateSaveBtn();',
        '}',
        'function renderCurrentView(){',
        '  grid.innerHTML="";var mObj=appData[modeEl.value];if(!mObj)return;',
        '  var wIdx=parseInt(weekEl.value,10)||0;',
        '  var weekObj=mObj.weeks[wIdx]||mObj.weeks.slice().shift()||{groups:{}};',
        '  var scheduleData=(weekObj.groups&&weekObj.groups[selectEl.value])||[];',
        '  updateSaveBtn();',
        '  scheduleData.forEach(function(d){',
        '    var col=document.createElement("div");col.className="day-column";',
        '    var h=document.createElement("div");h.className="day-header";h.textContent=d.day;col.appendChild(h);',
        '    var sb=document.createElement("div");sb.className="day-slots";',
        '    if(!d.slots.some(function(s){return s.subject;})){sb.innerHTML=\'<div class="day-off">🎉 Выходной</div>\';}else{',
        '      d.slots.forEach(function(s){',
        '        var c=document.createElement("div");',
        '        if(s.subject){',
        '          c.className="slot-card active-slot"+(s.moved_from?" moved-slot":"")+(s.custom_time?" shifted-slot":"");',
        '          var th=s.custom_time?\'<span class="old-time">\'+s.time+\'</span> <span class="custom-time">⏰ \'+s.custom_time+\'</span>\':\'<span class="pair-time">\'+s.time+\'</span>\';',
        '          c.innerHTML=\'<div class="slot-header"><span class="pair-num">\'+s.num+\' пара</span><div>\'+th+\'</div></div><div class="subject-title"></div><div class="slot-footer"></div>\';',
        '          c.querySelector(".subject-title").textContent=s.subject;',
        '          var f=c.querySelector(".slot-footer");',
        '          if(s.moved_from){var mb=document.createElement("span");mb.className="moved-badge";mb.textContent="🔄 с "+s.moved_from;f.appendChild(mb);}',
        '          if(s.link){let u=s.link;var b=document.createElement("button");b.className="online-btn";b.textContent="💻 Ссылка";b.onclick=function(){showLinkModal(u);};f.appendChild(b);}',
        '          if(s.room){var rb=document.createElement("span");rb.className="room-badge";rb.textContent="📍 "+s.room;f.appendChild(rb);}',
        '        }else{c.className="slot-card empty-slot";c.innerHTML=\'<div class="slot-header"><span class="pair-num">\'+s.num+\' пара</span><span class="pair-time">\'+s.time+\'</span></div><div class="empty-label">Свободно</div>\';}',
        '        sb.appendChild(c);',
        '      });',
        '    }',
        '    col.appendChild(sb);',
        '    if(d.removed&&d.removed.length>0){',
        '      var rs=document.createElement("div");rs.className="cancelled-section";rs.innerHTML=\'<div class="cancelled-title">🚫 Снято с расписания:</div>\';',
        '      d.removed.forEach(function(r){var it=document.createElement("div");it.className="cancelled-item";it.innerHTML=\'<span class="reason-tag">\'+r.reason+\'</span> <b>\'+r.num+\' пара:</b> <s>\'+r.subject+\'</s>\';rs.appendChild(it);});',
        '      col.appendChild(rs);',
        '    }',
        '    grid.appendChild(col);',
        '  });',
        '}',
        'modeEl.onchange=function(){populateMode(this.value,"");renderCurrentView();};',
        'selectEl.onchange=renderCurrentView;weekEl.onchange=renderCurrentView;',
        'function initApp(){',
        '  var saved=getSavedState(),defM=modeKeys.slice().shift(),tM=defM,tG="";',
        '  if(saved.indexOf("|")!==-1){var p=saved.split("|"),mC=p.slice().shift(),gC=p.pop();if(appData[mC]){tM=mC;tG=gC;}}',
        '  if(tM){modeEl.value=tM;populateMode(tM,tG);renderCurrentView();}',
        '}',
        'initApp();setTimeout(initApp,200);',
        'if("serviceWorker" in navigator){navigator.serviceWorker.register("sw.js");}',
        '</script></body></html>'
    ]
    with open(template_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
  

def get_val(sheet, row, col):
    cell = sheet.cell(row=row, column=col)
    if cell.value is not None:
        return str(cell.value).strip()
    for merged in sheet.merged_cells.ranges:
        if cell.coordinate in merged:
            val = sheet.cell(row=merged.min_row, column=merged.min_col).value
            return str(val).strip() if val is not None else ""
    return ""

def get_link_val(sheet, row, col):
    """
    Вытаскивает из ячейки и скрытую ссылку, и видимый текст (пароли, коды).
    Возвращает строку, где все данные склеены.
    """
    cell = sheet.cell(row=row, column=col)
    cell_text = str(cell.value or "").strip()
    
    # Если внутри ячейки зашита скрытая гиперссылка Excel
    if cell.hyperlink and cell.hyperlink.target:
        hyperlink_url = str(cell.hyperlink.target).strip()
        # Если текст поверх ссылки совпадает с ней или пустой, берем только ссылку
        if not cell_text or cell_text == hyperlink_url:
            return hyperlink_url
        # Если текст содержит пароль/информацию, склеиваем их через пробел
        return f"{hyperlink_url} (Информация: {cell_text})"
        
    return cell_text

def load_all_online_links(links_path):
    """
    Умный загрузчик ссылок: сам определяет формат файла.
    Для очников: C - Преподаватель, D - Ссылка, E - Пароль.
    Для ИУП:      C - Ссылка, D - Пароль.
    """
    all_links = []
    if not links_path or not os.path.exists(links_path):
        return all_links
    wb = openpyxl.load_workbook(links_path, data_only=True)
    sheet = wb.active
    
    # 🎯 АВТООПРЕДЕЛЕНИЕ ФОРМАТА: Ищем, в каком столбце лежит ссылка (в C или в D)
    # Проверяем первые 5 строк таблицы на наличие интернет-ссылки
    link_col = 4  # По умолчанию формат очников (столбец D)
    notes_col = 5 # По умолчанию формат очников (столбец E)
    
    for r in range(1, 6):
        val_c = str(sheet.cell(row=r, column=3).value or "").lower()
        val_d = str(sheet.cell(row=r, column=4).value or "").lower()
        # Если в столбце C написано 'ссылка' или текст начинается с http
        if "http" in val_c or "ссылк" in val_c or "zoom" in val_c or "teams" in val_c:
            link_col = 3   # Переключаемся на формат ИУП (ссылка в C)
            notes_col = 4  # Пароль в D
            break

    for r in range(1, sheet.max_row + 1):
        grp = get_val(sheet, r, 1)   # Столбец A - Группа
        subj = get_val(sheet, r, 2)  # Столбец B - Предмет
        
        # Читаем ссылку и пароль по динамически определенным столбцам
        raw_link = get_link_val(sheet, r, link_col)
        notes = get_val(sheet, r, notes_col)
        
        if grp and subj and raw_link:
            url_match = re.search(r'(https?://[^\s()]+)', raw_link)
            if url_match:
                clean_url = url_match.group(1)
                
                # Склеиваем ссылку с паролем, если он есть
                if notes:
                    final_payload = f"{clean_url} | {notes}"
                else:
                    final_payload = clean_url
                    
                all_links.append({
                    "group": grp.lower(),
                    "subject": " ".join(subj.lower().split()),
                    "link": final_payload
                })
    return all_links



def find_matching_link(group_name, subject_text, room_text, all_links):
    if not subject_text or not all_links:
        return None
    room_low = (room_text or "").lower()
    subj_low = subject_text.lower()
    online_markers = ("онлайн", "дист", "эиос", "online", "zoom", "сферум", "вебинар", "teams")
    has_online_word = any(m in room_low or m in subj_low for m in online_markers)
    has_digits_in_room = any(ch.isdigit() for ch in room_low)
    if has_digits_in_room and not has_online_word:
        return None

    grp_low = group_name.lower()
    norm_subject = " ".join(subj_low.split())
    best_link, best_len = None, 0
    for item in all_links:
        if (item["group"] in grp_low or grp_low in item["group"]) and item["subject"] in norm_subject:
            if len(item["subject"]) > best_len:
                best_link = item["link"]
                best_len = len(item["subject"])
    return best_link

def parse_subject_details(subject_text):
    res = {"subject": "", "custom_time": None, "is_removed": False, "remove_reason": None, "moved_from": None}
    if not subject_text:
        return res
    text = subject_text.strip()

    move_to = r"[\s\-–—(]*ПЕРЕНОС\s+НА\s+(\d{1,2}[./]\d{1,2}(?:[./]\d{2,4})?)[\s)!.]*$"
    m_to = re.search(move_to, text, flags=re.IGNORECASE)
    if m_to:
        res["is_removed"] = True
        res["remove_reason"] = f"➔ на {m_to.group(1)}"
        text = re.sub(move_to, "", text, flags=re.IGNORECASE).strip()

    move_from = r"[\s\-–—(]*ПЕРЕНОС\s+С\s+(\d{1,2}[./]\d{1,2}(?:[./]\d{2,4})?)[\s)!.]*$"
    m_from = re.search(move_from, text, flags=re.IGNORECASE)
    if m_from:
        res["moved_from"] = m_from.group(1)
        text = re.sub(move_from, "", text, flags=re.IGNORECASE).strip()

    cancel = r"[\s\-–—(]*ОТМЕНА[\s)!.]*$"
    if re.search(cancel, text, flags=re.IGNORECASE):
        res["is_removed"] = True
        res["remove_reason"] = "Отмена"
        text = re.sub(cancel, "", text, flags=re.IGNORECASE).strip()

    time_pat = r"^(?:[свСВ]\s*)?(\d{1,2}[:.]\d{2}(?:\s*[-–—]\s*\d{1,2}[:.]\d{2})?)\s*(.*)"
    m_time = re.match(time_pat, text, flags=re.DOTALL)
    if m_time:
        res["custom_time"] = m_time.group(1).replace(".", ":")
        text = m_time.group(2).strip(" -–—\n\t")

    res["subject"] = text
    return res

def parse_single_group(sheet, group_col, group_name, all_links):
    room_col = group_col + 1
    days_data, day_lookup, current_day = [], {}, None
    g_low = group_name.strip().lower()

    for r in range(2, sheet.max_row + 1):
        day_cell = get_val(sheet, r, 1)
        if day_cell:
            current_day = day_cell.replace("\n", " ")
            if current_day not in day_lookup:
                day_obj = {"day": current_day, "slots": [], "removed": []}
                days_data.append(day_obj)
                day_lookup[current_day] = day_obj

        if not current_day:
            continue

        pair_num = get_val(sheet, r, 2)
        pair_time = get_val(sheet, r, 3).replace("\n", " ")
        raw_subject = get_val(sheet, r, group_col)
        room = get_val(sheet, r, room_col).replace("\n", " ")

        # 1. ОТСЕКАЕМ ШАПКУ: Если это техническая строка с названиями столбцов
        if "номер" in pair_num.lower() or "время" in pair_time.lower() or "пара" in pair_num.lower():
            continue

        # 2. ФИЛЬТР ИУП: Если в предмете или кабинете содержится название группы
        if g_low in raw_subject.lower() or raw_subject.lower() in g_low:
            raw_subject = ""
        if g_low in room.lower() or room.lower() in g_low:
            room = ""

        # Дополнительная проверка на пустые технические строки
        if not pair_num and not pair_time and not raw_subject:
            continue

        info = parse_subject_details(raw_subject)
        online_link = find_matching_link(group_name, info["subject"], room, all_links)

        if pair_num or pair_time or raw_subject:
            if info["is_removed"]:
                day_lookup[current_day]["removed"].append({
                    "num": pair_num, 
                    "subject": info["subject"].replace("\n", " "), 
                    "reason": info["remove_reason"]
                })
                day_lookup[current_day]["slots"].append({
                    "num": pair_num, "time": pair_time, "custom_time": None, 
                    "subject": "", "room": "", "moved_from": None, "link": None
                })
            else:
                day_lookup[current_day]["slots"].append({
                    "num": pair_num, "time": pair_time, "custom_time": info["custom_time"], 
                    "subject": info["subject"], "room": room, "moved_from": info["moved_from"], "link": online_link
                })
    return days_data


def parse_workbook_all_groups(excel_path, all_links):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb.active
    groups_dict = {}
    skip_words = ("день", "пара", "время", "ауд", "каб", "место", "недел", "дата")
    col = 4
    while col <= sheet.max_column:
        title = get_val(sheet, 1, col).replace("\n", " ").strip()
        if title and not any(w in title.lower() for w in skip_words):
            groups_dict[title] = parse_single_group(sheet, col, title, all_links)
            col += 2
        else:
            col += 1
    return groups_dict


def generate_week_calendar(filename="schedule.xlsx", links_file="links.xlsx", output_html="index.html"):
    template_path = os.path.join(BASE_DIR, "template.html")
    output_path = os.path.join(BASE_DIR, output_html)
    meta_path = os.path.join(BASE_DIR, "weeks_meta.json")

    ensure_template(template_path)
    app_data = {}

    if os.path.exists(meta_path):
        with open(meta_path, "r", encoding="utf-8") as mf:
            raw_meta = json.load(mf)

        if isinstance(raw_meta, dict):
            for mode_key, mode_info in raw_meta.items():
                l_path = os.path.join(BASE_DIR, mode_info.get("links_file", links_file))
                if not os.path.exists(l_path): 
                    l_path = os.path.join(BASE_DIR, links_file)
                mode_links = load_all_online_links(l_path)

                weeks_list = []
                for w in mode_info.get("weeks", []):
                    f_path = os.path.join(BASE_DIR, w["file"])
                    if os.path.exists(f_path):
                        # Находим все группы в текущем файле без жесткого шага col += 2
                        wb = openpyxl.load_workbook(f_path, data_only=True)
                        sheet = wb.active
                        groups_dict = {}
                        skip_w = ("день", "пара", "время", "ауд", "каб", "место", "недел", "дата")
                        
                        for col in range(4, sheet.max_column + 1):
                            title = get_val(sheet, 1, col).replace("\n", " ").strip()
                            if title and not any(word in title.lower() for word in skip_w):
                                # Защита: если это имя группы, парсим её
                                groups_dict[title] = parse_single_group(sheet, col, title, mode_links)
                        
                        if groups_dict:
                            weeks_list.append({
                                "label": w["label"], 
                                "is_current": w.get("is_current", False), 
                                "groups": groups_dict
                            })
                if weeks_list:
                    app_data[mode_key] = {"title": mode_info.get("title", mode_key), "weeks": weeks_list}
    # Запасной вариант для локального одного файла, если meta_path не найден
    if not app_data:
        sched_path = os.path.join(BASE_DIR, filename)
        l_path = os.path.join(BASE_DIR, links_file)
        if os.path.exists(sched_path):
            mode_links = load_all_online_links(l_path)
            wb = openpyxl.load_workbook(sched_path, data_only=True)
            sheet = wb.active
            groups_dict = {}
            for col in range(4, sheet.max_column + 1):
                title = get_val(sheet, 1, col).replace("\n", " ").strip()
                if title and not any(w in title.lower() for w in ("день", "пара", "время", "ауд", "каб")):
                    groups_dict[title] = parse_single_group(sheet, col, title, mode_links)
            
            app_data["main"] = {
                "title": "🎓 Очное",
                "weeks": [{"label": "📅 Текущая неделя", "is_current": True, "groups": groups_dict}]
            }

    with open(template_path, "r", encoding="utf-8") as tf:
        html_content = tf.read()

    html_content = html_content.replace("__APP_DATA_JSON__", json.dumps(app_data, ensure_ascii=False))

    with open(output_path, "w", encoding="utf-8") as out:
        out.write(html_content)
    print(f"✅ Собрано потоков в календарь: {len(app_data)} -> {output_path}")

if __name__ == "__main__":
    generate_week_calendar("schedule.xlsx", "links.xlsx", "index.html")
